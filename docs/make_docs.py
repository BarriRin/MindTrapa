from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_heading_style(paragraph, size, bold=True, color=None):
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run(paragraph.text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0x1A, 0x53, 0x8B)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        run.bold = False
        run.italic = True
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    return p

def add_body(doc, text, indent=0):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    if indent:
        p.paragraph_format.left_indent = Inches(indent * 0.25)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    return p

def add_bullet(doc, text, indent=1):
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    run = p.add_run(text)
    run.font.size = Pt(11)
    p.paragraph_format.left_indent = Inches(indent * 0.3)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    return p

def add_separator(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run('─' * 72)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    return p

def add_status_row(doc, label, value):
    p = doc.add_paragraph()
    r1 = p.add_run(f'{label:<18}')
    r1.font.size = Pt(11)
    r1.bold = True
    r2 = p.add_run(value)
    r2.font.size = Pt(11)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    return p

# ─── RUSSIAN ────────────────────────────────────────────────────────────────

def build_ru(path):
    doc = Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # Title block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run('MINDTRAPA')
    r.font.size = Pt(28)
    r.bold = True
    r.font.color.rgb = RGBColor(0x1A, 0x53, 0x8B)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run('Описание проекта · Июль 2026')
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    sub.paragraph_format.space_after = Pt(14)

    add_separator(doc)

    # 1. Общая информация
    add_heading(doc, '1. Общая информация', 1)
    for label, val in [
        ('Название:', 'MindTrapa'),
        ('Жанр:', '3D головоломка-платформер'),
        ('Платформа:', 'Windows (Desktop)'),
        ('Стадия:', '1.0 (релиз)'),
        ('Разработчик:', 'Иван Данишевский (BarriRin), соло-проект'),
        ('Вдохновение:', 'levelDEVIL'),
    ]:
        add_status_row(doc, label, val)

    doc.add_paragraph()
    add_body(doc,
        'MindTrapa — это 3D-платформер, в котором игрок проходит уровни с ловушками, '
        'иллюзиями и нестандартной физикой. Каждый уровень — отдельная головоломка: '
        'мало просто прыгать, надо понять механику и найти безопасный путь '
        'к финишному флагу.')

    add_separator(doc)

    # 2. Технический стек
    add_heading(doc, '2. Технический стек', 1)
    for label, val in [
        ('Язык:', 'C++ (стандарт C++17)'),
        ('IDE:', 'Visual Studio 2022'),
        ('Графика:', 'DxLib 3.24f (обёртка над DirectX 11)'),
        ('Рендеринг:', '3D-модели (.mv1, .x), скайбокс, генерация звёздного неба'),
        ('Физика:', 'Собственный движок: AABB-коллизии, deltaTime-физика (60 FPS)'),
        ('Анимации:', 'Стейт-машина: Idle, Run, Jump, Fall, Death, Wave/Yes (победа)'),
        ('Ресурсы:', 'Quaternius, Kenney, KayKit (все CC0); музыка — Suno AI; SFX — Freesound.org (CC0)'),
        ('Локализация:', 'EN / RU / JA, отдельные шрифты под кириллицу и кану'),
        ('Сохранения:', '5 слотов профилей, INI-формат: звёзды, рекорды, настройки, паки оформления'),
    ]:
        add_status_row(doc, label, val)

    add_separator(doc)

    # 3. Концепция
    add_heading(doc, '3. Концепция игры', 1)
    add_body(doc,
        'Игрок управляет персонажем в 3D-пространстве, перемещаясь по платформам '
        'к цели (флаг-триггер) на каждом уровне. Главная идея — каждый блок '
        '(группа из 10 уровней) вводит одну-две новые механики и строит '
        'на них всю прогрессию сложности.')
    doc.add_paragraph()
    add_heading(doc, 'Ключевые принципы дизайна:', 2)
    add_bullet(doc, 'Честная сложность: каждая смерть — следствие решения игрока, не случайности.')
    add_bullet(doc, 'Читаемость: опасность должна угадываться визуально или через первую попытку.')
    add_bullet(doc, 'Прогрессия: уровни внутри блока плавно усложняют одну механику без резких перепадов.')

    add_separator(doc)

    # 4. Механики
    add_heading(doc, '4. Механики игры (17 уникальных)', 1)

    groups = [
        ('Базовые платформы', [
            ('PLATFORM', 'обычная твёрдая платформа'),
            ('INVISIBLE_WALL', 'невидимая стена (только коллизия)'),
            ('FAKE_PLATFORM', 'выглядит как платформа, но игрок проваливается'),
            ('FAKE_SPIKES', 'выглядит как шипы, но безопасна'),
            ('TRIGGER', 'финиш уровня (флаг)'),
        ]),
        ('Опасности', [
            ('SPIKES', 'мгновенная смерть при касании'),
            ('RETRACTABLE_SPIKES', 'шипы, выдвигающиеся по таймеру'),
            ('PENDULUM_BLADE', 'маятник-пила, качается из стороны в сторону'),
        ]),
        ('Динамические платформы', [
            ('MOVING', 'движется между двумя точками (ping-pong)'),
            ('DISAPPEARING', 'исчезает через секунду после наступания'),
            ('CRUMBLING', 'рассыпается под ногами при длительном стоянии'),
        ]),
        ('Специальные', [
            ('BUTTON', 'активирует / деактивирует другие блоки'),
            ('TELEPORT', 'мгновенно перемещает в заданную точку'),
        ]),
        ('Физические зоны', [
            ('GRAVITY_ZONE', 'меняет гравитацию: пониженная, повышенная или обратная (ходьба по потолку)'),
            ('LIGHT_PULSE_ZONE', 'периодическая вспышка света в тёмных уровнях Блока 4'),
            ('ICE_PLATFORM', 'скользкая поверхность: инерция, прыжок недоступен'),
            ('BOUNCE_PAD', 'батут: единственный способ прыгнуть со льда'),
        ]),
    ]

    for group_name, items in groups:
        add_heading(doc, group_name, 2)
        for name, desc in items:
            p = doc.add_paragraph()
            r_name = p.add_run(f'{name:<22}')
            r_name.bold = True
            r_name.font.size = Pt(10)
            r_name.font.color.rgb = RGBColor(0x1A, 0x53, 0x8B)
            r_desc = p.add_run(desc)
            r_desc.font.size = Pt(11)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)

    add_separator(doc)

    # 5. Уровни
    add_heading(doc, '5. Структура уровней (5 блоков × 10 уровней = 50 уровней)', 1)

    blocks_data = [
        ('Блок 1 — Основы (уровни 1–10)',
         'Знакомство с базовыми механиками',
         'платформы, шипы, невидимые стены, фейковые платформы, телепорты, кнопки, '
         'движущиеся и исчезающие платформы',
         'Завершён ✓'),
        ('Блок 2 — Тайминг и точность (уровни 11–20)',
         'Уровни требуют точного выбора момента',
         'выдвижные шипы, рассыпающиеся платформы, фейковые шипы',
         'Завершён ✓'),
        ('Блок 3 — Гравитация и маятники (уровни 21–30)',
         'Нестандартная физика',
         'зоны гравитации (пониженная / повышенная), маятники-пилы',
         'Завершён ✓'),
        ('Блок 4 — Темнота (уровни 31–40)',
         'Ограниченная видимость',
         'накопительная темнота при движении, зоны мерцания',
         'Завершён ✓'),
        ('Блок 5 — Лёд и батуты (уровни 41–50)',
         'Инерция и альтернативные прыжки',
         'ледяные платформы (скользкость, запрет прыжка), батуты',
         'Завершён ✓'),
    ]

    for title_b, theme, mechanics, status in blocks_data:
        add_heading(doc, title_b, 2)
        for lbl, val in [('Тема:', theme), ('Механики:', mechanics), ('Статус:', status)]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r1 = p.add_run(f'{lbl:<12}')
            r1.bold = True
            r1.font.size = Pt(11)
            r2 = p.add_run(val)
            r2.font.size = Pt(11)

    add_separator(doc)

    # 6. Текущее состояние
    add_heading(doc, '6. Текущее состояние (июль 2026 — версия 1.0)', 1)

    add_heading(doc, 'Реализовано:', 2)
    done = [
        '50 из 50 уровней — завершены, пройдены и отбалансированы автором',
        'Все 17 механик реализованы и протестированы',
        'Полный pipeline 3D-моделей: собственный конвертер OBJ→X, темы под каждый блок',
        'Персонаж игрока с полной стейт-машиной анимаций (Idle/Run/Jump/Fall/Death/победа)',
        'Сферический скайбокс с процедурным фоллбеком, туман по биомам, мерцающие звёзды',
        'Полное меню: выбор блока/уровня, настройки с живым применением слайдеров, пауза',
        'Система профилей: 5 слотов, звёзды, рекорды, постоянное сохранение на диск',
        'Разблокируемые паки оформления (скайбокс + музыка) за суммарные звёзды',
        'Полная локализация: EN / RU / JA, включая отдельные шрифты под кириллицу и кану',
        'Стриминг фоновой музыки по блокам + SFX (прыжок, приземление, смерть, победа, батут, телепорт)',
        'Независимый аудит кода (3 агента, 36 находок) — все дофикшены и проверены',
    ]
    for item in done:
        add_bullet(doc, item)

    add_separator(doc)

    # 7. Возможные направления развития
    add_heading(doc, '7. Возможные направления развития (пост-релиз)', 1)
    add_body(doc,
        'Проект завершён и закрыт как версия 1.0. Ниже — идеи, которые не входили '
        'в исходный скоуп и не планируются в ближайшее время, но могли бы стать '
        'отдельным дополнением:')
    doc.add_paragraph()
    future = [
        'Режим «Один забег» (Hardcore Run): фиксированное количество жизней на все 50 уровней, единый таймер, отдельная таблица рекордов',
        'Локальная таблица рекордов / speedrun-режим',
        'Экран достижений с косметическими разблокировками (скины персонажа, доп. модели блоков)',
        'Возможный портинг на другие платформы',
    ]
    for item in future:
        add_bullet(doc, item)

    add_separator(doc)

    # Footer
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer_p.add_run('van4ous00@gmail.com  ·  GitHub: BarriRin / MindTrapa')
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.save(path)
    print(f'Saved: {path}')



# ─── ENGLISH ────────────────────────────────────────────────────────────────

def build_en(path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run('MINDTRAPA')
    r.font.size = Pt(28); r.bold = True
    r.font.color.rgb = RGBColor(0x1A, 0x53, 0x8B)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run('Project Overview · July 2026')
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    sub.paragraph_format.space_after = Pt(14)

    add_separator(doc)

    add_heading(doc, '1. General Information', 1)
    for label, val in [
        ('Title:', 'MindTrapa'),
        ('Genre:', '3D Puzzle Platformer'),
        ('Platform:', 'Windows (Desktop)'),
        ('Stage:', '1.0 (Release)'),
        ('Developer:', 'Ivan Danisheuski (BarriRin), solo project'),
        ('Inspiration:', 'levelDEVIL'),
    ]:
        add_status_row(doc, label, val)

    doc.add_paragraph()
    add_body(doc,
        'MindTrapa is a 3D platformer in which the player navigates levels filled with '
        'traps, illusions, and unconventional physics. Each level is a standalone puzzle: '
        'simply jumping is not enough — the player must understand the mechanics and '
        'find a safe path to the finish flag.')

    add_separator(doc)

    add_heading(doc, '2. Tech Stack', 1)
    for label, val in [
        ('Language:', 'C++ (C++17)'),
        ('IDE:', 'Visual Studio 2022'),
        ('Graphics:', 'DxLib 3.24f (DirectX 11 wrapper)'),
        ('Rendering:', '3D models (.mv1, .x), skybox, procedural star field'),
        ('Physics:', 'Custom engine: AABB collisions, deltaTime physics (60 FPS)'),
        ('Animations:', 'State machine: Idle, Run, Jump, Fall, Death, Wave/Yes (celebration)'),
        ('Assets:', 'Quaternius, Kenney, KayKit (all CC0); music — Suno AI; SFX — Freesound.org (CC0)'),
        ('Localization:', 'EN / RU / JA, separate font sets for Cyrillic and Kana'),
        ('Save system:', '5 profile slots, INI format: stars, best times, settings, theme packs'),
    ]:
        add_status_row(doc, label, val)

    add_separator(doc)

    add_heading(doc, '3. Game Concept', 1)
    add_body(doc,
        'The player controls a character in 3D space, moving across platforms toward '
        'a goal (flag trigger) on each level. The core idea: each block (a group of '
        '10 levels) introduces one or two new mechanics and builds its entire difficulty '
        'progression around them.')
    doc.add_paragraph()
    add_heading(doc, 'Core Design Principles:', 2)
    add_bullet(doc, 'Fair difficulty: every death is the result of a player decision, not random chance.')
    add_bullet(doc, 'Readability: danger should be readable visually or through a first attempt — never through luck.')
    add_bullet(doc, 'Progression: levels within a block gradually escalate one mechanic without sudden difficulty spikes.')

    add_separator(doc)

    add_heading(doc, '4. Game Mechanics (17 unique)', 1)

    groups = [
        ('Basic Platforms', [
            ('PLATFORM', 'standard solid platform'),
            ('INVISIBLE_WALL', 'invisible wall (collision only)'),
            ('FAKE_PLATFORM', 'looks like a platform, but the player falls through'),
            ('FAKE_SPIKES', 'looks like spikes, but is safe'),
            ('TRIGGER', 'level finish (flag)'),
        ]),
        ('Hazards', [
            ('SPIKES', 'instant death on contact'),
            ('RETRACTABLE_SPIKES', 'spikes that extend and retract on a timer'),
            ('PENDULUM_BLADE', 'swinging saw blade pendulum'),
        ]),
        ('Dynamic Platforms', [
            ('MOVING', 'moves between two points (ping-pong)'),
            ('DISAPPEARING', 'vanishes one second after being stepped on'),
            ('CRUMBLING', 'breaks apart after standing on it too long'),
        ]),
        ('Special', [
            ('BUTTON', 'activates / deactivates other blocks'),
            ('TELEPORT', 'instantly moves the player to a set location'),
        ]),
        ('Physics Zones', [
            ('GRAVITY_ZONE', 'alters gravity: low, high, or reversed (walk on the ceiling)'),
            ('LIGHT_PULSE_ZONE', 'periodic light flash inside Block 4\'s darkness-affected levels'),
            ('ICE_PLATFORM', 'slippery surface: inertia applies, jumping is disabled'),
            ('BOUNCE_PAD', 'launch pad: the only way to jump from ice'),
        ]),
    ]

    for group_name, items in groups:
        add_heading(doc, group_name, 2)
        for name, desc in items:
            p = doc.add_paragraph()
            r_name = p.add_run(f'{name:<22}')
            r_name.bold = True; r_name.font.size = Pt(10)
            r_name.font.color.rgb = RGBColor(0x1A, 0x53, 0x8B)
            r_desc = p.add_run(desc)
            r_desc.font.size = Pt(11)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)

    add_separator(doc)

    add_heading(doc, '5. Level Structure (5 blocks × 10 levels = 50 levels)', 1)

    blocks_data = [
        ('Block 1 — Basics (Levels 1–10)',
         'Introduction to core mechanics',
         'platforms, spikes, invisible walls, fake platforms, teleports, buttons, moving and disappearing platforms',
         'Complete ✓'),
        ('Block 2 — Timing & Precision (Levels 11–20)',
         'Levels require precise timing',
         'retractable spikes, crumbling platforms, fake spikes',
         'Complete ✓'),
        ('Block 3 — Gravity & Blades (Levels 21–30)',
         'Non-standard physics',
         'gravity zones (low / high gravity), pendulum blades',
         'Complete ✓'),
        ('Block 4 — Darkness (Levels 31–40)',
         'Limited visibility',
         'cumulative darkness while moving, flickering light zones',
         'Complete ✓'),
        ('Block 5 — Ice & Bounce (Levels 41–50)',
         'Inertia and alternative jumping',
         'ice platforms (slipperiness, no jumping), bounce pads',
         'Complete ✓'),
    ]

    for title_b, theme, mechanics, status in blocks_data:
        add_heading(doc, title_b, 2)
        for lbl, val in [('Theme:', theme), ('Mechanics:', mechanics), ('Status:', status)]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r1 = p.add_run(f'{lbl:<12}')
            r1.bold = True; r1.font.size = Pt(11)
            r2 = p.add_run(val); r2.font.size = Pt(11)

    add_separator(doc)

    add_heading(doc, '6. Current State (July 2026 — version 1.0)', 1)
    add_heading(doc, 'Implemented:', 2)
    for item in [
        '50 of 50 levels — complete, played through and balanced by the author',
        'All 17 mechanics implemented and tested',
        'Full 3D model pipeline: custom OBJ→X converter, per-block themes',
        'Player character with a full animation state machine (Idle/Run/Jump/Fall/Death/celebration)',
        'Spherical skybox with procedural fallback, per-biome fog, twinkling starfield',
        'Full menu system: block/level select, settings with live slider apply, pause',
        'Profile system: 5 slots, stars, best times, persistent save to disk',
        'Unlockable cosmetic theme packs (skybox + music), unlocked by total star count',
        'Full localization: EN / RU / JA, including dedicated font sets for Cyrillic and Kana',
        'Streamed per-block background music + SFX (jump, land, death, win, bounce, teleport)',
        'Independent code audit (3 agents, 36 findings) — all fixed and verified',
    ]:
        add_bullet(doc, item)

    add_separator(doc)

    add_heading(doc, '7. Possible Future Directions (Post-Release)', 1)
    add_body(doc,
        'The project is complete and closed out as version 1.0. The ideas below were '
        'not part of the original scope and are not currently planned, but could form '
        'a separate future addition:')
    doc.add_paragraph()
    for item in [
        '"One Run" (Hardcore) mode: fixed lives across all 50 levels, single global timer, separate leaderboard',
        'Local leaderboard / speedrun mode',
        'Achievements screen with cosmetic unlocks (character skins, extra block models)',
        'Potential porting to other platforms',
    ]:
        add_bullet(doc, item)

    add_separator(doc)

    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer_p.add_run('van4ous00@gmail.com  ·  GitHub: BarriRin / MindTrapa')
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.save(path)
    print(f'Saved: {path}')


# ─── JAPANESE ───────────────────────────────────────────────────────────────

def build_ja(path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run('MINDTRAPA')
    r.font.size = Pt(28); r.bold = True
    r.font.color.rgb = RGBColor(0x1A, 0x53, 0x8B)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run('プロジェクト概要 · 2026年7月')
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    sub.paragraph_format.space_after = Pt(14)

    add_separator(doc)

    add_heading(doc, '1. 基本情報', 1)
    for label, val in [
        ('タイトル:', 'MindTrapa'),
        ('ジャンル:', '3Dパズルプラットフォーマー'),
        ('プラットフォーム:', 'Windows（デスクトップ）'),
        ('開発段階:', '1.0（リリース版）'),
        ('開発者:', 'イワン・ダニシェフスキー（BarriRin）、個人開発'),
        ('インスピレーション:', 'levelDEVIL'),
    ]:
        add_status_row(doc, label, val)

    doc.add_paragraph()
    add_body(doc,
        'MindTrapaは、罠・錯覚・独自物理法則が組み合わさった3Dプラットフォーマーです。'
        '各レベルはひとつの独立したパズルとして設計されており、'
        'ただジャンプするだけでなく、仕組みを理解してゴールのフラグまでの安全なルートを見つける必要があります。')

    add_separator(doc)

    add_heading(doc, '2. 技術スタック', 1)
    for label, val in [
        ('言語:', 'C++（C++17規格）'),
        ('IDE:', 'Visual Studio 2022'),
        ('グラフィック:', 'DxLib 3.24f（DirectX 11ラッパー）'),
        ('レンダリング:', '3Dモデル（.mv1, .x）、スカイボックス、星空生成'),
        ('物理演算:', '自作エンジン：AABB衝突判定、deltaTime物理（60FPS基準）'),
        ('アニメーション:', 'ステートマシン：待機・走り・ジャンプ・落下・死亡・勝利'),
        ('アセット:', 'Quaternius, Kenney, KayKit（すべてCC0）；音楽はSuno AI、効果音はFreesound.org（CC0）'),
        ('ローカライズ:', 'EN／RU／JA、キリル文字・かな用に専用フォントを使用'),
        ('セーブ:', 'プロファイル5枠、INI形式：スター・ベストタイム・設定・テーマパック'),
    ]:
        add_status_row(doc, label, val)

    add_separator(doc)

    add_heading(doc, '3. ゲームコンセプト', 1)
    add_body(doc,
        'プレイヤーは3D空間でキャラクターを操作し、各レベルのゴール（フラグトリガー）を目指します。'
        'コアコンセプトは「各ブロック（10レベルのグループ）が1〜2種類の新メカニクスを導入し、'
        'その難易度進行をそれだけで構成する」というものです。')
    doc.add_paragraph()
    add_heading(doc, 'レベルデザインの基本方針:', 2)
    add_bullet(doc, '公平な難易度：すべての死はプレイヤーの判断の結果であり、理不尽な偶然によるものではない。')
    add_bullet(doc, '視認性：危険はビジュアルまたは最初の試みで読み取れるべきであり、運任せにしない。')
    add_bullet(doc, '段階的な難易度上昇：ブロック内のレベルは一つのメカニクスを段階的に難しくし、急激な難易度変化を避ける。')

    add_separator(doc)

    add_heading(doc, '4. ゲームメカニクス（全17種類）', 1)

    groups = [
        ('基本プラットフォーム', [
            ('PLATFORM', '通常の固い床'),
            ('INVISIBLE_WALL', '透明な壁（衝突判定のみ）'),
            ('FAKE_PLATFORM', '床に見えるが、踏むと落ちる'),
            ('FAKE_SPIKES', 'トゲに見えるが、安全'),
            ('TRIGGER', 'レベルのゴール（フラグ）'),
        ]),
        ('危険物', [
            ('SPIKES', '触れると即死'),
            ('RETRACTABLE_SPIKES', 'タイマーで出し入れするトゲ'),
            ('PENDULUM_BLADE', '左右に揺れる刃の振り子'),
        ]),
        ('動的プラットフォーム', [
            ('MOVING', '2点間を往復する床'),
            ('DISAPPEARING', '乗ってから1秒後に消える床'),
            ('CRUMBLING', '長く乗ると崩れる床'),
        ]),
        ('特殊ブロック', [
            ('BUTTON', '他のブロックを有効化／無効化するボタン'),
            ('TELEPORT', '指定地点に瞬間移動'),
        ]),
        ('物理ゾーン', [
            ('GRAVITY_ZONE', '重力を変化させるゾーン：低重力・高重力・逆重力（天井を歩く）'),
            ('LIGHT_PULSE_ZONE', 'ブロック4の暗闇エリアで周期的に光るゾーン'),
            ('ICE_PLATFORM', '滑りやすい床：慣性あり、ジャンプ不可'),
            ('BOUNCE_PAD', 'バウンスパッド：氷の上から唯一飛び出せる手段'),
        ]),
    ]

    for group_name, items in groups:
        add_heading(doc, group_name, 2)
        for name, desc in items:
            p = doc.add_paragraph()
            r_name = p.add_run(f'{name:<22}')
            r_name.bold = True; r_name.font.size = Pt(10)
            r_name.font.color.rgb = RGBColor(0x1A, 0x53, 0x8B)
            r_desc = p.add_run(desc); r_desc.font.size = Pt(11)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)

    add_separator(doc)

    add_heading(doc, '5. レベル構成（5ブロック × 10レベル = 全50レベル）', 1)

    blocks_data = [
        ('ブロック1 — 基礎（レベル1〜10）',
         '基本メカニクスの習得',
         '床・トゲ・透明壁・フェイク床・テレポート・ボタン・動く床・消える床',
         '完成 ✓'),
        ('ブロック2 — タイミングと精度（レベル11〜20）',
         '正確なタイミングが要求されるレベル群',
         '引き込み式トゲ・崩れる床・フェイクトゲ',
         '完成 ✓'),
        ('ブロック3 — 重力と刃（レベル21〜30）',
         '独自物理法則の活用',
         '重力ゾーン（低重力・高重力）、振り子ブレード',
         '完成 ✓'),
        ('ブロック4 — 暗闇（レベル31〜40）',
         '視界制限のある環境',
         '移動するほど暗くなる累積暗闇、明滅ゾーン',
         '完成 ✓'),
        ('ブロック5 — 氷とバウンス（レベル41〜50）',
         '慣性と代替ジャンプ手段',
         '氷床（滑り・ジャンプ不可）、バウンスパッド',
         '完成 ✓'),
    ]

    for title_b, theme, mechanics, status in blocks_data:
        add_heading(doc, title_b, 2)
        for lbl, val in [('テーマ:', theme), ('メカニクス:', mechanics), ('ステータス:', status)]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r1 = p.add_run(f'{lbl}  ')
            r1.bold = True; r1.font.size = Pt(11)
            r2 = p.add_run(val); r2.font.size = Pt(11)

    add_separator(doc)

    add_heading(doc, '6. 現在の開発状況（2026年7月 — バージョン1.0）', 1)
    add_heading(doc, '実装済み:', 2)
    for item in [
        '50レベル中50レベル — 全て完成、作者本人がプレイ・調整済み',
        '全17メカニクス — 実装・テスト済み',
        '完全な3Dモデルパイプライン：自作OBJ→Xコンバーター、ブロックごとのテーマ',
        'プレイヤーキャラクターに完全なアニメーションステートマシン（待機・走り・ジャンプ・落下・死亡・勝利）',
        '球面スカイボックス（プロシージャルフォールバック付き）、バイオームごとの霧、瞬くスターフィールド',
        '完全なメニューシステム：ブロック・レベル選択、ライブ反映される設定スライダー、ポーズ',
        'プロファイルシステム：5枠、スター、ベストタイム、ディスクへの永続保存',
        '解放可能なテーマパック（スカイボックス＋音楽）、累計スター数で解放',
        '完全なローカライズ：EN／RU／JA、キリル文字・かな用の専用フォント込み',
        'ブロックごとのストリーミングBGM＋効果音（ジャンプ・着地・死亡・勝利・バウンス・テレポート）',
        '独立コードレビュー（3エージェント、36件の指摘）— 全て修正・検証済み',
    ]:
        add_bullet(doc, item)

    add_separator(doc)

    add_heading(doc, '7. 今後の可能性（リリース後）', 1)
    add_body(doc,
        '本プロジェクトはバージョン1.0として完成・クローズしています。以下は当初のスコープには'
        '含まれておらず、現時点で計画もされていませんが、将来的な追加コンテンツとして'
        '考えられるアイデアです：')
    doc.add_paragraph()

    for item in [
        '「ワンラン」モード（Hardcore Run）：全50レベルで残機制・単一グローバルタイマー・専用ランキング',
        'ローカルランキング／スピードランモード',
        '実績画面と外見解放（キャラクタースキン、追加3Dモデル）',
        '他プラットフォームへの移植検討',
    ]:
        add_bullet(doc, item)

    add_separator(doc)

    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer_p.add_run('van4ous00@gmail.com  ·  GitHub: BarriRin / MindTrapa')
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.save(path)
    print(f'Saved: {path}')


build_ru(r'D:\Probuem\MindTrapa\docs\MindTrapa_RU.docx')
build_en(r'D:\Probuem\MindTrapa\docs\MindTrapa_EN.docx')
build_ja(r'D:\Probuem\MindTrapa\docs\MindTrapa_JA.docx')
